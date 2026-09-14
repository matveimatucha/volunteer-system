#!/usr/bin/env python3
"""Parse the legacy volunteer xlsx into JSON for import-legacy-volunteers.js."""
import hashlib
import json
import re
import sys

import openpyxl

MONTHS = {
    'январ': 1, 'феврал': 2, 'март': 3, 'апрел': 4,
    'мая': 5, 'май': 5, 'июн': 6, 'июл': 7,
    'август': 8, 'сентябр': 9, 'октябр': 10, 'ноябр': 11, 'декабр': 12
}
MONTH_LAST_DAY = [0, 31, 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]


def month_num(word):
    w = (word or '').lower()
    for prefix, num in MONTHS.items():
        if prefix in w:
            return num
    return None


def iso(year, month, day):
    return f'{year:04d}-{month:02d}-{day:02d}'


def parse_date_range(raw, year):
    if not raw:
        return '', ''
    text = str(raw).strip().lower().replace('ё', 'е')
    text = re.sub(r'\s+', ' ', text)
    text = text.replace('(ночь)', '').replace('с ', '').replace(' на ', '-')
    text = text.replace('–', '-').replace('—', '-')

    # февраль-апрель / март-апрель
    m = re.fullmatch(r'([а-я]+)\s*-\s*([а-я]+)', text)
    if m and month_num(m.group(1)) and month_num(m.group(2)):
        a, b = month_num(m.group(1)), month_num(m.group(2))
        return iso(year, a, 1), iso(year, b, MONTH_LAST_DAY[b])

    # 1 апреля-6 мая / 16-25 марта / 9-11 июля / 6 апреля- 6 мая
    m = re.search(
        r'(\d{1,2})(?:\s+([а-я]+))?\s*-\s*(\d{1,2})\s+([а-я]+)',
        text
    )
    if m:
        d1, mon1, d2, mon2 = m.group(1), m.group(2), m.group(3), m.group(4)
        b = month_num(mon2)
        a = month_num(mon1) if mon1 else b
        if a and b:
            return iso(year, a, int(d1)), iso(year, b, int(d2))

    # 23 января / 8 февраля
    m = re.search(r'(\d{1,2})\s+([а-я]+)', text)
    if m and month_num(m.group(2)):
        a = month_num(m.group(2))
        day = int(m.group(1))
        return iso(year, a, day), iso(year, a, day)

    return '', ''


def ru_date(iso_s):
    if not iso_s:
        return ''
    y, m, d = iso_s.split('-')
    months = [
        '', 'января', 'февраля', 'марта', 'апреля', 'мая', 'июня',
        'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря'
    ]
    return f'{int(d)} {months[int(m)]} {y}'


def slug_person(last, first, middle, vk):
    base = (vk or f'{last}|{first}|{middle}').lower().strip()
    return hashlib.sha1(base.encode('utf-8')).hexdigest()[:12]


def marked(value):
    if value is None or value == '':
        return False
    if isinstance(value, (int, float)) and value != 0:
        return True
    return str(value).strip() not in ('', '0', 'нет', '-', '—')


def parse(path, year=2026):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    dates = [c.value for c in ws[2]]
    titles = [c.value for c in ws[3]]

    events = []
    for idx in range(6, ws.max_column):
        title = titles[idx] if idx < len(titles) else None
        if not title or not str(title).strip():
            continue
        date_raw, date_end = parse_date_range(dates[idx] if idx < len(dates) else None, year)
        n = len(events) + 1
        events.append({
            'id': f'legacy-{year}-{n:02d}',
            'col': idx,
            'title': str(title).strip(),
            'dateLabel': str(dates[idx]).strip() if dates[idx] else '',
            'dateRaw': date_raw,
            'dateEndRaw': date_end if date_end and date_end != date_raw else '',
            'date': ru_date(date_raw),
            'dateEnd': ru_date(date_end) if date_end and date_end != date_raw else ''
        })

    people = []
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, values_only=True):
        last = str(row[0] or '').strip()
        first = str(row[1] or '').strip()
        if not last and not first:
            continue
        middle = str(row[2] or '').strip()
        faculty = str(row[3] or '').strip()
        vk = str(row[4] or '').strip()
        attended = []
        for ev in events:
            cell = row[ev['col']] if ev['col'] < len(row) else None
            if marked(cell):
                attended.append(ev['id'])
        people.append({
            'id': slug_person(last, first, middle, vk),
            'lastName': last,
            'firstName': first,
            'middleName': middle,
            'faculty': faculty,
            'vk': vk,
            'attended': attended
        })

    return {'year': year, 'events': events, 'people': people}


def main():
    src = sys.argv[1]
    year = int(sys.argv[2]) if len(sys.argv) > 2 else 2026
    dest = sys.argv[3] if len(sys.argv) > 3 else None
    data = parse(src, year)
    text = json.dumps(data, ensure_ascii=False, indent=2)
    if dest:
        with open(dest, 'w', encoding='utf-8') as f:
            f.write(text)
    else:
        sys.stdout.reconfigure(encoding='utf-8')
        print(text)


if __name__ == '__main__':
    main()
